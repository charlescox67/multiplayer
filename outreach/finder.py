#!/usr/bin/env python3
"""
Pipeline 1 -- The Finder (see outreach-pipeline-context.md).

Builds outreach/leads.xlsx: one tab of ICP-matching leads for the human approval
gate, one tab of community "connectors" who aren't direct leads. Charles reviews
the Leads tab and sets Status to Approved/Rejected by hand -- nothing here sends
anything or scrapes anything live.

There is no product yet. This is pre-build research: the draft messages ask how
a person's team actually uses AI tools and whether there's any shared visibility
across teammates, not a pitch to test something. Both tabs are sorted by
confidence (high to low) so the most-verified people surface first.

This first run is seeded from outreach/leads_seed.py: rows found via manual web
research this session (see that file's docstring for the verification approach and
its limits). There is intentionally no automated LinkedIn/Reddit scraper wired in --
outreach-pipeline-context.md flags that unofficial LinkedIn automation risks account
bans and that only official APIs or Charles's own logged-in tooling should ever touch
LinkedIn directly. To extend sourcing, either:
  - append more manually-verified rows to leads_seed.py, or
  - point SOURCE_FN at a real data source (e.g. a Sales Navigator / Apollo export
    loaded from CSV) that returns dicts shaped like the ones in leads_seed.py.

Usage:
    pip install -r outreach/requirements.txt
    python3 outreach/finder.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from leads_seed import CONNECTORS, LEADS  # noqa: E402

COLUMNS = [
    ("name", "Name", 22),
    ("title", "Title", 34),
    ("company", "Company", 38),
    ("profile_link", "Profile / Post Link", 46),
    ("contact_method", "Best Contact Method", 26),
    ("confidence", "Confidence", 12),
    ("summary", "Profile Summary", 55),
    ("match_reason", "Match Reason", 50),
    ("subject", "Subject", 40),
    ("draft_message", "Draft Message", 65),
    ("status", "Status", 14),
    ("date_found", "Date Found", 14),
]

CONFIDENCE_ORDER = {"high": 0, "medium": 1, "low": 2}
_LOW_SIGNALS = (
    "unconfirmed", "not confirmed", "single search snippet",
    "lower-confidence", "lower confidence", "caveat",
    "self-described", "self described",
)
_HIGH_SIGNALS = ("confirmed",)


def confidence(row: dict) -> str:
    """Derived from the caveats already written into summary/match_reason --
    not a separate judgment call, just surfacing what's already on the record."""
    text = (row.get("summary", "") + " " + row.get("match_reason", "")).lower()
    if any(s in text for s in _LOW_SIGNALS):
        return "low"
    if any(s in text for s in _HIGH_SIGNALS):
        return "high"
    return "medium"

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
LINK_FONT = Font(color="1155CC", underline="single")

# Personalization style per outreach-pipeline-context.md: simple, specific, no
# overselling, consistent core ask with a varying reference point. There is no
# working "click to send" URL for LinkedIn DMs (not supported), so this is a
# copy-paste draft, not an auto-send link -- click Profile Link, click Message,
# paste, edit, send.
#
# No product exists yet, and the message doesn't say otherwise -- no "I'm
# building," no "before building anything," no mention of multiplayer or any
# product. Just the personalized detail and direct questions about how the
# person's team actually uses AI and whether that usage is shared or solo.
OPERATOR_TEMPLATE = (
    "Hey {first_name}, saw {personal_detail}. Question for you: does your "
    "team use ChatGPT or Claude much, and if so, is it mostly everyone working "
    "solo, or is there any shared visibility into what people are trying or "
    "deciding? Curious how you all handle it."
)
CONSULTANT_TEMPLATE = (
    "Hey {first_name}, saw {personal_detail}. Question for you: across the "
    "teams or clients you work with, is AI tool usage mostly a solo thing per "
    "person, or is there any shared visibility into what's being tried? "
    "Curious what you're seeing."
)
CONNECTOR_TEMPLATE = (
    "Hey {first_name}, saw {personal_detail}. Question for you: is solo AI "
    "tool usage with little shared visibility across a team something you "
    "hear about often from your community? Curious what you're seeing."
)

# Only LinkedIn InMail (messaging someone outside your network) has a subject
# field; a message to an existing 1st-degree connection has none, so this is
# blank-safe to ignore there. Same rule as the message body: no product name,
# no "building," just what the note is about.
SUBJECT_BY_KIND = {
    "operator": "Question about your team and AI tools",
    "consultant": "Question about your clients and AI tools",
}
CONNECTOR_SUBJECT = "Question about AI usage in your community"


def subject(row: dict) -> str:
    return SUBJECT_BY_KIND.get(row.get("kind"), CONNECTOR_SUBJECT)


def draft_message(row: dict) -> str:
    first_name = row["name"].split()[0]
    template = {
        "operator": OPERATOR_TEMPLATE,
        "consultant": CONSULTANT_TEMPLATE,
    }.get(row.get("kind"), CONNECTOR_TEMPLATE)
    return template.format(first_name=first_name, personal_detail=row.get("personal_detail", ""))


def dedupe(rows: list[dict]) -> list[dict]:
    seen = set()
    out = []
    for row in rows:
        key = row["profile_link"].strip().lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(row)
    return out


def sort_by_confidence(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: CONFIDENCE_ORDER[confidence(r)])


def write_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    for col_idx, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    ordered_rows = sort_by_confidence(dedupe(rows))
    for row_idx, row in enumerate(ordered_rows, start=2):
        for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            if key == "draft_message":
                value = draft_message(row)
            elif key == "subject":
                value = subject(row)
            elif key == "confidence":
                value = confidence(row)
            else:
                value = row.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = WRAP
            if key == "profile_link" and row.get(key):
                cell.hyperlink = row[key]
                cell.font = LINK_FONT

    if rows:
        status_col = next(i for i, (k, _, _) in enumerate(COLUMNS, start=1) if k == "status")
        dv = DataValidation(type="list", formula1='"Approved,Rejected"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}{len(rows) + 1}")


def build(output_path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "Leads", LEADS)
    write_sheet(wb, "Connectors", CONNECTORS)
    wb.save(output_path)
    print(f"Wrote {len(dedupe(LEADS))} lead(s) and {len(dedupe(CONNECTORS))} connector(s) to {output_path}")


if __name__ == "__main__":
    build(Path(__file__).parent / "leads.xlsx")
